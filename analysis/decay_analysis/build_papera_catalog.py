# -*- coding: utf-8 -*-
"""扫描 PaperA 归档目录，写出案例总表.xlsx。"""
from __future__ import annotations

import csv
import json
import os
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

_d = os.path.dirname(os.path.abspath(__file__))
while True:
    if os.path.isfile(os.path.join(_d, "README.md")):
        if _d not in sys.path:
            sys.path.insert(0, _d)
        break
    _parent = os.path.dirname(_d)
    if _parent == _d:
        raise RuntimeError("Cannot find wellbore_moc_method root")
    _d = _parent

from moc_simulate.config import FRACTURE_CONFIG

ROOT = Path(r"e:\water_hammer_research\wellbore_moc_method")
OUTPUT = ROOT / "output"
PAPERA = OUTPUT / "PaperA井口多裂缝水击响应"
XLSX = PAPERA / "案例总表.xlsx"

DEFAULT_CF = float(FRACTURE_CONFIG["Cf"])
DEFAULT_KLEAK = float(FRACTURE_CONFIG["kleak"])

CASE_N = {
    "single": 1,
    "dual": 2,
    "triple": 3,
    "quad": 4,
    "quint": 5,
    "hex": 6,
    "hept": 7,
    "oct": 8,
}

HEADERS = [
    "案例编号",
    "来源",
    "摩阻",
    "初缝x1_m",
    "间距D_m",
    "缝数n",
    "Cf",
    "kleak",
    "Brunone_k",
    "波形相对路径",
    "独立正演",
    "备注",
]


def rel(p: Path) -> str:
    try:
        return str(p.relative_to(PAPERA)).replace("\\", "/")
    except ValueError:
        return str(p)


def collect() -> list[dict]:
    rows: list[dict] = []

    decay_csv = PAPERA / "01_几何网格" / "峰值表" / "decay_table.csv"
    decay_npz_dir = PAPERA / "01_几何网格" / "波形_npz"
    table_keys = set()
    with decay_csv.open(encoding="utf-8") as f:
        for r in csv.DictReader(f):
            if int(r["frac_idx"]) != 1:
                continue
            fr = r["friction_model"]
            x1 = float(r["x1"])
            d = float(r["spacing_m"])
            n = int(r["n_total"])
            table_keys.add((fr, int(x1), int(d), n))
            npz = decay_npz_dir / f"{fr}_x1_{int(x1)}_sp_{int(d)}_n_{n}.npz"
            csv_wave = (
                PAPERA
                / "01_几何网格"
                / "波形"
                / f"{fr}_x1_{int(x1)}_sp_{int(d)}_n_{n}"
                / "moc_timeseries.csv"
            )
            wave = csv_wave if csv_wave.is_file() else npz
            independent = "是" if n > 1 or int(d) == 10 else "否（复用同x1单缝波形）"
            note = ""
            if n == 1:
                note = "单缝几何与D无关；每个(摩阻,x1)只正演一次，本行是D格子副本" if int(d) != 10 else "单缝独立正演（canonical D=10，其余D复制此波形）"
            rows.append(
                {
                    "来源": "衰减几何网格",
                    "摩阻": fr,
                    "初缝x1_m": x1,
                    "间距D_m": d,
                    "缝数n": n,
                    "Cf": DEFAULT_CF,
                    "kleak": DEFAULT_KLEAK,
                    "Brunone_k": "" if fr == "steady" else "Vardy k(Re)",
                    "波形相对路径": rel(wave) if wave.is_file() else "",
                    "独立正演": independent,
                    "备注": note,
                }
            )

    pat = re.compile(r"(steady|brunone)_x1_(\d+)_sp_(\d+)_n_(\d+)\.npz")
    for p in sorted(decay_npz_dir.glob("*.npz")):
        m = pat.match(p.name)
        if not m:
            continue
        fr, x1, d, n = m.group(1), int(m.group(2)), int(m.group(3)), int(m.group(4))
        if (fr, x1, d, n) in table_keys:
            continue
        csv_wave = (
            PAPERA
            / "01_几何网格"
            / "波形"
            / f"{fr}_x1_{x1}_sp_{d}_n_{n}"
            / "moc_timeseries.csv"
        )
        wave = csv_wave if csv_wave.is_file() else p
        rows.append(
            {
                "来源": "衰减几何网格",
                "摩阻": fr,
                "初缝x1_m": float(x1),
                "间距D_m": float(d),
                "缝数n": n,
                "Cf": DEFAULT_CF,
                "kleak": DEFAULT_KLEAK,
                "Brunone_k": "" if fr == "steady" else "Vardy k(Re)",
                "波形相对路径": rel(wave),
                "独立正演": "是",
                "备注": "仅有npz，未写入decay_table.csv（x1=1000 的 n=2–8）",
            }
        )

    cf_csv = PAPERA / "02_裂缝属性_CfKleak" / "cf_kleak_table.csv"
    wave_dir = PAPERA / "02_裂缝属性_CfKleak" / "波形_npz"
    wave_csv_root = PAPERA / "02_裂缝属性_CfKleak" / "波形"
    if cf_csv.is_file():
        seen = set()
        with cf_csv.open(encoding="utf-8") as f:
            for r in csv.DictReader(f):
                if int(r["frac_idx"]) != 1:
                    continue
                key = (r["friction_model"], r["Cf"], r["Kleak"], r["n_total"])
                if key in seen:
                    continue
                seen.add(key)
                fr = r["friction_model"]
                cf = float(r["Cf"])
                kl = float(r["Kleak"])
                npz_name = f"{fr}_cf_{cf:.1e}_kleak_{kl:.1e}.npz"
                npz = wave_dir / npz_name
                csv_wave = wave_csv_root / f"{fr}_cf_{cf:.1e}_kleak_{kl:.1e}" / "moc_timeseries.csv"
                wave = csv_wave if csv_wave.is_file() else npz
                rows.append(
                    {
                        "来源": "Cf×kleak扫描",
                        "摩阻": fr,
                        "初缝x1_m": float(r["x1"]),
                        "间距D_m": float(r["spacing_m"]),
                        "缝数n": int(float(r["n_total"])),
                        "Cf": cf,
                        "kleak": kl,
                        "Brunone_k": "" if fr == "steady" else "Vardy k(Re)",
                        "波形相对路径": rel(wave) if wave.is_file() else "",
                        "独立正演": "是",
                        "备注": "该扫描的n在3与5之间不统一，见本行缝数n",
                    }
                )

    leak_root = PAPERA / "03_leakoff验证"
    for js in sorted(leak_root.rglob("moc_leakoff.json")):
        djson = json.loads(js.read_text(encoding="utf-8"))
        cfg = djson.get("config") or {}
        x_f = cfg.get("x_f") or []
        case = js.parent.name
        fr_key = js.parent.parent.name
        n = CASE_N.get(case, len(x_f) if x_f else "")
        x1 = float(x_f[0]) if x_f else ""
        if len(x_f) >= 2:
            spacing = float(x_f[1]) - float(x_f[0])
        else:
            m = re.search(r"_D(\d+)$", fr_key)
            spacing = float(m.group(1)) if m else (50.0 if case == "single" else "")
        csv_p = js.with_name("moc_timeseries.csv")
        fric = cfg.get("friction") or ("brunone" if fr_key.startswith("brunone") else "steady")
        rows.append(
            {
                "来源": "leakoff验证",
                "摩阻": fric,
                "初缝x1_m": x1,
                "间距D_m": spacing,
                "缝数n": n,
                "Cf": float(cfg.get("Cf", DEFAULT_CF)),
                "kleak": float(cfg.get("kleak", DEFAULT_KLEAK)),
                "Brunone_k": "" if fric == "steady" else "Vardy k(Re)",
                "波形相对路径": rel(csv_p if csv_p.is_file() else js),
                "独立正演": "是",
                "备注": f"目录 {fr_key}/{case}",
            }
        )

    et = PAPERA / "04_能量回归" / "energy_table.csv"
    if et.is_file():
        with et.open(encoding="utf-8") as f:
            for r in csv.DictReader(f):
                fr = r["friction"]
                case = r["case"]
                d = float(r["spacing_m"])
                n = int(r["n_fracs"])
                cf = float(r["Cf"])
                kl = float(r["kleak"])
                geo_dir = PAPERA / "04_能量回归" / "工况" / fr / f"{case}_D{int(d)}"
                prop_dir = (
                    PAPERA
                    / "04_能量回归"
                    / "工况"
                    / fr
                    / f"{case}_cf{cf:.1e}_kl{kl:.1e}"
                )
                is_default = abs(cf - DEFAULT_CF) < 1e-18 and abs(kl - DEFAULT_KLEAK) < 1e-18
                if is_default and geo_dir.is_dir():
                    folder = geo_dir
                elif prop_dir.is_dir():
                    folder = prop_dir
                elif geo_dir.is_dir():
                    folder = geo_dir
                else:
                    folder = prop_dir
                ej = folder / "energy.json"
                csv_wave = folder / "moc_timeseries.csv"
                wave = csv_wave if csv_wave.is_file() else ej
                rows.append(
                    {
                        "来源": "能量回归",
                        "摩阻": fr,
                        "初缝x1_m": 4100.0,
                        "间距D_m": d,
                        "缝数n": n,
                        "Cf": float(r["Cf"]),
                        "kleak": float(r["kleak"]),
                        "Brunone_k": "" if fr == "steady" else "Vardy k(Re)",
                        "波形相对路径": rel(wave) if wave.is_file() else rel(folder),
                        "独立正演": "是",
                        "备注": "初缝按 leakoff 默认 4100 m",
                    }
                )

    bs = PAPERA / "05_Brunone常数k" / "工况"
    kpat = re.compile(r"^D(\d+)_k([0-9.]+)$")
    for ddir in sorted(bs.iterdir()):
        if not ddir.is_dir():
            continue
        m = kpat.match(ddir.name)
        if not m:
            continue
        csv_p = ddir / "moc_timeseries.csv"
        if not csv_p.is_file():
            continue
        k = float(m.group(2))
        rows.append(
            {
                "来源": "Brunone常数k×间距",
                "摩阻": "steady" if k == 0.0 else "brunone",
                "初缝x1_m": 4100.0,
                "间距D_m": float(m.group(1)),
                "缝数n": 4,
                "Cf": DEFAULT_CF,
                "kleak": DEFAULT_KLEAK,
                "Brunone_k": k,
                "波形相对路径": rel(csv_p),
                "独立正演": "是",
                "备注": "k为脚本钉死的常数，不是 Vardy k(Re)",
            }
        )

    return rows


def write_xlsx(records: list[dict]) -> None:
    wb = Workbook()
    font = Font(name="Arial", size=10)
    font_h = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    font_title = Font(name="Arial", size=14, bold=True)
    fill_h = PatternFill("solid", fgColor="0F4D92")
    fill_note = PatternFill("solid", fgColor="FFF2CC")
    thin = Border(
        left=Side(style="thin", color="CFCECE"),
        right=Side(style="thin", color="CFCECE"),
        top=Side(style="thin", color="CFCECE"),
        bottom=Side(style="thin", color="CFCECE"),
    )
    wrap = Alignment(wrap_text=True, vertical="center")

    # 说明
    ws0 = wb.active
    ws0.title = "说明"
    ws0["A1"] = "PaperA 井口多裂缝水击响应 — 案例归档说明"
    ws0["A1"].font = font_title
    notes = [
        ("归档日期", "2026-08-14"),
        ("归档目录", str(PAPERA)),
        ("目录结构", "01几何网格 / 02 CfKleak / 03 leakoff / 04能量回归 / 05 Brunone常数k；见同目录「目录说明.txt」"),
        ("纳入", "decay 几何网格、Cf×kleak、leakoff、能量回归、Brunone 常数k×间距"),
        ("未纳入", "LHS 随机集、分层基准、反演/神经算子/可辨识性结果"),
        ("单缝补全", "steady/brunone × x1={2000,2500,3000,3500,4000,4500} × D=10…100，共120格"),
        ("单缝正演", "单缝与D无关：只做了12次独立MOC，其余108个D格子复制同一波形npz"),
        ("decay_table", "补全后 960 工况（原840的n=2–8 + 120的n=1）；CSV 4320行"),
        ("x1=1000", "n=2–8 的npz仍在 decay_regression/01_simulated_waves，未写入decay_table"),
        ("Cf×kleak", "几何固定 x1=3000 m、D=20 m；缝数n在3与5之间不统一"),
        ("默认Cf/kleak", f"Cf={DEFAULT_CF:g} m²，kleak={DEFAULT_KLEAK:g}（几何网格/leakoff/k扫描）"),
    ]
    ws0["A3"] = "项目"
    ws0["B3"] = "内容"
    ws0["A3"].font = font_h
    ws0["B3"].font = font_h
    ws0["A3"].fill = fill_h
    ws0["B3"].fill = fill_h
    for i, (k, v) in enumerate(notes, start=4):
        ws0[f"A{i}"] = k
        ws0[f"B{i}"] = v
        ws0[f"A{i}"].font = Font(name="Arial", size=10, bold=True)
        ws0[f"B{i}"].font = font
        ws0[f"B{i}"].alignment = wrap
        if "单缝" in k:
            ws0[f"A{i}"].fill = fill_note
            ws0[f"B{i}"].fill = fill_note
    ws0.column_dimensions["A"].width = 18
    ws0.column_dimensions["B"].width = 92
    ws0.row_dimensions[1].height = 22
    for i in range(4, 14):
        ws0.row_dimensions[i].height = 28

    # 全部案例
    ws = wb.create_sheet("全部案例")
    for col, h in enumerate(HEADERS, 1):
        cell = ws.cell(1, col, h)
        cell.font = font_h
        cell.fill = fill_h
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for i, rec in enumerate(records, start=2):
        vals = [
            i - 1,
            rec["来源"],
            rec["摩阻"],
            rec["初缝x1_m"],
            rec["间距D_m"],
            rec["缝数n"],
            rec["Cf"],
            rec["kleak"],
            rec["Brunone_k"],
            rec["波形相对路径"],
            rec["独立正演"],
            rec["备注"],
        ]
        for col, v in enumerate(vals, 1):
            cell = ws.cell(i, col, v)
            cell.font = font
            cell.border = thin
            cell.alignment = Alignment(vertical="center")
            if col in (7, 8) and isinstance(v, float):
                cell.number_format = "0.00E+00"
            if col in (4, 5) and isinstance(v, float):
                cell.number_format = "0"
    last = 1 + len(records)
    tab = Table(displayName="Cases", ref=f"A1:L{last}")
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(tab)
    widths = [12, 20, 12, 14, 12, 10, 12, 12, 16, 62, 16, 56]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.auto_filter.ref = f"A1:L{last}"
    ws.freeze_panes = "A2"

    # 总览（公式）
    ws1 = wb.create_sheet("总览", 1)
    ws1["A1"] = "案例计数（公式，随全部案例表更新）"
    ws1["A1"].font = font_title
    ws1.merge_cells("A1:C1")
    headers1 = ["统计项", "公式计数", "说明"]
    for col, h in enumerate(headers1, 1):
        c = ws1.cell(3, col, h)
        c.font = font_h
        c.fill = fill_h
    items = [
        ("全部案例行数", f"=COUNTA(Cases[案例编号])", "一行=一个(来源,摩阻,x1,D,n,Cf,kleak[,k])工况"),
        ("衰减几何网格", f'=COUNTIF(Cases[来源],"衰减几何网格")', "含 n=1 补全与 x1=1000 仅npz"),
        ("其中单缝 n=1", f'=COUNTIFS(Cases[来源],"衰减几何网格",Cases[缝数n],1)', "120 格；独立正演仅12次"),
        ("单缝独立正演", f'=COUNTIFS(Cases[来源],"衰减几何网格",Cases[缝数n],1,Cases[独立正演],"是")', "应为12"),
        ("Cf×kleak扫描", f'=COUNTIF(Cases[来源],"Cf×kleak扫描")', "2摩阻×8 Cf×9 kleak"),
        ("leakoff验证", f'=COUNTIF(Cases[来源],"leakoff验证")', "x1=4100 m"),
        ("能量回归", f'=COUNTIF(Cases[来源],"能量回归")', "energy_table.csv"),
        ("Brunone常数k×间距", f'=COUNTIF(Cases[来源],"Brunone常数k×间距")', "5个D × 6个k"),
        ("steady 行数", f'=COUNTIF(Cases[摩阻],"steady")', ""),
        ("brunone 行数", f'=COUNTIF(Cases[摩阻],"brunone")', ""),
    ]
    for i, (name, formula, note) in enumerate(items, start=4):
        ws1.cell(i, 1, name).font = font
        cell = ws1.cell(i, 2, formula)
        cell.font = Font(name="Arial", size=10, color="000000")
        cell.alignment = Alignment(horizontal="center")
        ws1.cell(i, 3, note).font = font
    ws1.column_dimensions["A"].width = 28
    ws1.column_dimensions["B"].width = 16
    ws1.column_dimensions["C"].width = 48

    # 单缝补全清单
    ws2 = wb.create_sheet("单缝补全")
    ws2["A1"] = "本次新跑的12次独立单缝正演（D=10 canonical；其余D为npz复制）"
    ws2["A1"].font = font_title
    ws2.merge_cells("A1:G1")
    h2 = ["摩阻", "初缝x1_m", "canonical间距_m", "缝数n", "Cf", "kleak", "npz"]
    for col, h in enumerate(h2, 1):
        c = ws2.cell(3, col, h)
        c.font = font_h
        c.fill = fill_h
    irow = 4
    for rec in records:
        if rec["来源"] != "衰减几何网格":
            continue
        if rec["缝数n"] != 1:
            continue
        if rec["独立正演"] != "是":
            continue
        vals = [rec["摩阻"], rec["初缝x1_m"], rec["间距D_m"], 1, rec["Cf"], rec["kleak"], rec["波形相对路径"]]
        for col, v in enumerate(vals, 1):
            cell = ws2.cell(irow, col, v)
            cell.font = font
            if col in (5, 6) and isinstance(v, float):
                cell.number_format = "0.00E+00"
        irow += 1
    ws2.cell(irow + 1, 1, "独立正演次数")
    ws2.cell(irow + 1, 2, f"=COUNTA(B4:B{irow-1})")
    ws2.cell(irow + 1, 2).font = Font(name="Arial", size=10, bold=True)
    for i, w in enumerate([12, 14, 18, 10, 12, 12, 70], 1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    wb.save(XLSX)


def main() -> None:
    if not PAPERA.is_dir():
        raise FileNotFoundError(PAPERA)
    records = collect()
    write_xlsx(records)
    print(f"records={len(records)}")
    print(f"xlsx={XLSX}")


if __name__ == "__main__":
    main()
